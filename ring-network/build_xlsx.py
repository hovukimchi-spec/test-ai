"""Dựng ring_network_NhaBe.xlsx từ model.build(). Chạy: python3 build_xlsx.py"""
import os,sys,itertools
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
import model
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule,FormulaRule

d=model.build();C=d['config']
HDR=PatternFill('solid',fgColor='1F3864');HF=Font(bold=True,color='FFFFFF')
RING_COL={'R1':'9DC3E6','R2':'F4B183','R3':'A9D18E','R4':'C9B3E6','R5':'FFD966','R6':'8FD9D9','R7':'D7B99B','R8':'F7B6D2'}
GOOD=PatternFill('solid',fgColor='C6EFCE');BAD=PatternFill('solid',fgColor='FFC7CE');WARN=PatternFill('solid',fgColor='FFEB9C')
GREY=PatternFill('solid',fgColor='EDEDED');INP=PatternFill('solid',fgColor='FFF2CC')
thin=Side(style='thin',color='BFBFBF');BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
wb=Workbook()
def sheet(name,first=False):
    ws=wb.active if first else wb.create_sheet();ws.title=name;ws.sheet_view.showGridLines=False;return ws
def header(ws,row,cols,widths=None):
    for i,c in enumerate(cols,1):
        x=ws.cell(row,i,c);x.fill=HDR;x.font=HF;x.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center');x.border=BORDER
    if widths:
        for i,w in enumerate(widths,1):ws.column_dimensions[CL(i)].width=w
    ws.row_dimensions[row].height=32
def title(ws,text,sub=None):
    ws['A1']=text;ws['A1'].font=Font(bold=True,size=15,color='1F3864')
    if sub:ws['A2']=sub;ws['A2'].font=Font(italic=True,color='7F7F7F')
def box(ws,r0,r1,c0,c1):
    for r in range(r0,r1+1):
        for c in range(c0,c1+1):ws.cell(r,c).border=BORDER

# ---------------- Tham số ----------------
P=sheet('Tham số')
title(P,'Tham số quy ước (sửa ô vàng, mọi công thức tự cập nhật)')
rows=[('Ngưỡng suy hao đoạn < 10 km (dB)',12.1),('Ngưỡng suy hao đoạn 10–40 km (dB)',17),('Ngưỡng suy hao đoạn > 40 km (dB)',27),
      ('Mốc chiều dài 1 (km)',10),('Mốc chiều dài 2 (km)',40),('Chênh suy hao tối đa giữa 2 core (dB)',C['core_delta_max']),
      ('Hệ số cáp dự trữ ước tính (km thực = km KML × hệ số)',C['slack_est']),('Cảnh báo OTDR/KML nhỏ hơn',C['ratio_warn'][0]),
      ('Cảnh báo OTDR/KML lớn hơn',C['ratio_warn'][1]),('Cảnh báo lệch OTDR so với hợp đồng (tỷ lệ)',C['hd_warn'])]
header(P,3,['Tham số','Giá trị'],[62,14])
for i,(k,v) in enumerate(rows,4):
    P.cell(i,1,k);x=P.cell(i,2,v);x.fill=INP
box(P,4,3+len(rows),1,2)
P['A16']='Quy ước: suy hao tính theo TỪNG đoạn giữa hai trạm liền kề (mỗi trạm có SFP thu/phát riêng), không cộng dồn cả ring. Ngưỡng ứng với SFP 10/40/80 km.'
P['A17']='Nguồn dữ liệu: file ghi tay của người dùng (BBNT, ODF, MX, độ dài cáp, KML). BBNT chỉ là con số thống nhất giữa hai bên.'
T1,T2,T3,M1,M2,CD,SL,RL,RH,HW=[f"'Tham số'!$B${i}" for i in range(4,14)]

# ---------------- Đoạn ----------------
S=sheet('Đoạn & quỹ suy hao')
title(S,'Đoạn cáp giữa hai trạm liền kề — quỹ suy hao theo từng đoạn','Ô trắng = dữ liệu từ BBNT/KML; cột công thức tự tính. Đoạn đổi tuyến (HNB066, HNB005): dùng km KML × hệ số dự trữ.')
cols=['Mã đoạn','Ring','Từ','Đến','Loại','Lề đường','km HĐ (BBNT)','km OTDR (BBNT)','km KML','Đổi tuyến','km tính','dB core #1','dB core #2','dB lớn nhất','Ngưỡng (dB)','Dư (dB)','Chênh core (dB)','Trạng thái','OTDR/KML','HĐ lệch OTDR','Mã BBNT']
header(S,4,cols,[10,6,10,10,10,16,10,10,9,9,9,9,9,9,9,9,9,11,9,10,8])
segs=sorted(d['segments'],key=lambda s:(s['ring'],s['seq']))
for i,s in enumerate(segs,5):
    vals=[s['id'],s['ring'],s['a'],s['z'],s['kind'],s['roadside'],s['km_hd'],s['km_otdr'],s['km_kml'],'Có' if s['rerouted'] else '']
    for j,v in enumerate(vals,1):S.cell(i,j,v)
    S.cell(i,11,f'=IF(J{i}="Có",I{i}*{SL},IF(H{i}="",I{i}*{SL},H{i}))')
    S.cell(i,12,s['db1']);S.cell(i,13,s['db2'])
    S.cell(i,14,f'=IF(L{i}="","",MAX(L{i},M{i})*IF(J{i}="Có",K{i}/H{i},1))')
    S.cell(i,15,f'=IF(K{i}<{M1},{T1},IF(K{i}<={M2},{T2},{T3}))')
    S.cell(i,16,f'=IF(N{i}="","",O{i}-N{i})')
    S.cell(i,17,f'=IF(L{i}="","",ABS(L{i}-M{i}))')
    S.cell(i,18,f'=IF(N{i}="","Chưa đo",IF(AND(N{i}<=O{i},Q{i}<={CD}),"Đạt","Cảnh báo"))')
    S.cell(i,19,f'=IF(OR(H{i}="",I{i}=""),"",H{i}/I{i})')
    S.cell(i,20,f'=IF(OR(G{i}="",H{i}=""),"",ABS(H{i}-G{i})/G{i})')
    S.cell(i,21,s['bbnt'])
    for c in (11,14,16,17,19):S.cell(i,c).number_format='0.00'
    S.cell(i,20).number_format='0%'
    S.cell(i,1).fill=PatternFill('solid',fgColor=RING_COL[s['ring']])
last=4+len(segs);box(S,5,last,1,21);S.freeze_panes='E5';S.auto_filter.ref=f'A4:U{last}'
S.conditional_formatting.add(f'R5:R{last}',CellIsRule(operator='equal',formula=['"Đạt"'],fill=GOOD))
S.conditional_formatting.add(f'R5:R{last}',CellIsRule(operator='equal',formula=['"Cảnh báo"'],fill=BAD))
S.conditional_formatting.add(f'R5:R{last}',CellIsRule(operator='equal',formula=['"Chưa đo"'],fill=WARN))
S.conditional_formatting.add(f'S5:S{last}',FormulaRule(formula=[f'AND(S5<>"",OR(S5<{RL},S5>{RH}))'],fill=WARN))
S.conditional_formatting.add(f'T5:T{last}',FormulaRule(formula=[f'AND(T5<>"",T5>{HW})'],fill=WARN))
SR=f"'Đoạn & quỹ suy hao'!"
rng=lambda col:f"{SR}${col}$5:${col}${last}"

# ---------------- Ring ----------------
R=sheet('Ring')
title(R,'8 ring AGG-HCMNBE12 ↔ AGG-HCMNBE22','Tổng km chỉ để tham khảo; quỹ suy hao KHÔNG cộng dồn.')
header(R,4,['Ring','Chuỗi (từ NBE12 sang NBE22)','Số CSG','Số đoạn ring','Tổng km tính','Đoạn dài nhất (km)','Dư thấp nhất (dB)','Đoạn cảnh báo','Nhánh cụt'],[7,62,8,9,11,12,12,10,22])
for i,(rid,r) in enumerate(d['rings'].items(),5):
    chain=' ─ '.join(['NBE12']+r['csg']+['NBE22']);sp=[s['z'] for s in d['segments'] if s['ring']==rid and s['kind']!='ring']
    R.cell(i,1,rid).fill=PatternFill('solid',fgColor=RING_COL[rid]);R.cell(i,2,chain);R.cell(i,3,len(r['csg']))
    R.cell(i,4,f'=COUNTIFS({rng("B")},A{i},{rng("E")},"ring")')
    R.cell(i,5,f'=SUMIFS({rng("K")},{rng("B")},A{i},{rng("E")},"ring")').number_format='0.0'
    R.cell(i,6,f'=_xlfn.MAXIFS({rng("K")},{rng("B")},A{i},{rng("E")},"ring")').number_format='0.0'
    R.cell(i,7,f'=_xlfn.MINIFS({rng("P")},{rng("B")},A{i},{rng("E")},"ring")').number_format='0.0'
    R.cell(i,8,f'=COUNTIFS({rng("B")},A{i},{rng("R")},"Cảnh báo")');R.cell(i,9,', '.join(sp) or '—')
box(R,5,12,1,9)
R.conditional_formatting.add('H5:H12',CellIsRule(operator='greaterThan',formula=['0'],fill=WARN))
# sơ đồ dạng ô
R['A15']='Sơ đồ (số trên mũi tên = km tính; ■ AGG, ○ CSG)';R['A15'].font=Font(bold=True)
row=16
for rid,r in d['rings'].items():
    chain=['NBE12']+r['csg']+['NBE22'];rs=sorted([s for s in segs if s['ring']==rid and s['kind']=='ring'],key=lambda s:s['seq'])
    R.cell(row,1,rid).fill=PatternFill('solid',fgColor=RING_COL[rid]);R.cell(row,1).font=Font(bold=True)
    c=2
    for k,n in enumerate(chain):
        x=R.cell(row,c,('■ ' if n.startswith('NBE') else '○ ')+n);x.font=Font(bold=True,color='FFFFFF' if n.startswith('NBE') else '000000')
        x.fill=PatternFill('solid',fgColor='000000' if n.startswith('NBE') else RING_COL[rid]);x.alignment=Alignment(horizontal='center');c+=1
        if k<len(rs):
            srow=5+segs.index(rs[k]);y=R.cell(row,c,f"={SR}K{srow}");y.number_format='0.0" km ►"';y.alignment=Alignment(horizontal='center');y.font=Font(color='7F7F7F')
            if rs[k]['roadside']:R.cell(row+1,c,'lề '+rs[k]['roadside'].split(' ')[0]).font=Font(size=8,italic=True,color='7F7F7F');R.cell(row+1,c).alignment=Alignment(horizontal='center')
            c+=1
    row+=2
for col in range(3,20):
    if R.column_dimensions[CL(col)].width is None or col>9:R.column_dimensions[CL(col)].width=11

# ---------------- Nút ----------------
N=sheet('Nút')
title(N,'Nút: trạm cấp 1 (AGG) và CSG','Toạ độ theo KML. Nút nhánh cụt (cấp 3) chỉ có một hướng đến.')
header(N,4,['Mã','Loại','Ring','Vai trò','Vĩ độ','Kinh độ','Bản đồ'],[10,8,7,22,11,11,14])
nl=sorted(d['nodes'].values(),key=lambda n:(n['type']!='AGG',n.get('ring',''),n['id']))
for i,n in enumerate(nl,5):
    for j,v in enumerate([n['id'],n['type'],n.get('ring',''),n.get('role','trạm cấp 1 (AGG)' if n['type']=='AGG' else ''),n['lat'],n['lon']],1):N.cell(i,j,v)
    N.cell(i,7,f'=HYPERLINK("https://maps.google.com/?q="&E{i}&","&F{i},"Mở bản đồ")').font=Font(color='0563C1',underline='single')
box(N,5,4+len(nl),1,7);NL=4+len(nl);N.freeze_panes='A5'

# ---------------- ODF CSG ----------------
O=sheet('ODF CSG')
title(O,'ODF tại các CSG — lưới cổng (hướng đến)','Mỗi ODF 12 cổng (nhánh cụt 6). Màu theo hướng. Chi tiết đầu cáp ở sheet "ODF chi tiết". Ô trống = chưa ghi trong file nguồn.')
header(O,4,['ODF']+[f'Cổng {i}' for i in range(1,13)]+['Cổng có dữ liệu','Hướng 1','Hướng 2'],[12]+[10]*12+[11,12,12])
pal={}
def color(dest):
    if dest not in pal:pal[dest]=['DDEBF7','FCE4D6','E2EFDA','FFF2CC','E4DFEC','D9D9D9','F8CBAD','C6E0B4'][len(pal)%8]
    return PatternFill('solid',fgColor=pal[dest])
nodes_odf=list(dict.fromkeys(x['node'] for x in d['odf']))
for i,nd in enumerate(nodes_odf,5):
    O.cell(i,1,'ODF-'+nd).font=Font(bold=True)
    ps=[x for x in d['odf'] if x['node']==nd];dests=[]
    for x in ps:
        if x['port']>12:continue
        c=O.cell(i,1+x['port'],x['dest'].replace('AGG-','') if x['dest'] else '');c.alignment=Alignment(horizontal='center')
        if x['dest']:c.fill=color(x['dest']);dests.append(x['dest'])
    O.cell(i,14,f'=COUNTA(B{i}:M{i})-COUNTBLANK(B{i}:M{i})+COUNTBLANK(B{i}:M{i})-COUNTBLANK(B{i}:M{i})' if False else f'=12-COUNTBLANK(B{i}:M{i})' if len(ps)>6 else f'=6-COUNTBLANK(B{i}:G{i})')
    u=list(dict.fromkeys(dests));O.cell(i,15,u[0] if u else '');O.cell(i,16,u[1] if len(u)>1 else '')
box(O,5,4+len(nodes_odf),1,16);O.freeze_panes='B5'
OD=sheet('ODF chi tiết')
title(OD,'ODF CSG — chi tiết từng cổng (trích nguyên văn ghi chú)')
header(OD,3,['Nút','Cổng','Nhãn E','Hướng đến','Số FO (ghi chú)','Chỉ số mét (cs)','Ghi chú gốc'],[10,7,8,14,10,12,80])
for i,x in enumerate(d['odf'],4):
    for j,v in enumerate([x['node'],x['port'],x['label'],x['dest'],x['fo'],x['cs'],x['cable']],1):OD.cell(i,j,v)
OD.auto_filter.ref=f'A3:G{3+len(d["odf"])}';OD.freeze_panes='A4'
OA=sheet('ODF AGG')
title(OA,'ODF tại hai trạm AGG (NBE22 trái / Nguyễn Bình–Nguyễn Văn Tạo phải)','Các cột số là giá trị OTDR/chỉ số ghi tay (mét), ngày ghi trong file nguồn 10–11/2020; ý nghĩa từng cột chưa được xác nhận.')
header(OA,4,['Phía','ODF','Cổng','Cột O','Đích ghi chú','Số 1 (m)','Số 2 (m)','Số 3 (m)','Ghi chú','Cột T','Cột U'],[7,34,6,7,22,9,9,9,40,9,9])
for i,x in enumerate(d['odf_agg'],5):
    for j,k in enumerate(['phia','odf','cong','cot_O','dich','cot1','cot2','cot3','ghi_chu','cot_T','cot_U'],1):
        v=x[k];OA.cell(i,j,int(v) if str(v).isdigit() else v)
OA.freeze_panes='A5'

# ---------------- Cáp ----------------
CB=sheet('Cáp')
title(CB,'Danh sách cáp theo sổ "độ dài cáp Nhà Bè"','Đầu/cuối là măng xông hoặc trạm; độ dài mét; cột cs = chỉ số mét ở hai đầu cáp (nhiều ô trống: cần bổ sung để đổi OTDR ↔ vị trí).')
header(CB,4,["STT",'Đầu','Cuối','cs đầu','cs cuối','Mã cáp','Số FO','Dài (m)','Ghi chú'],[5,28,28,9,9,22,8,10,50])
for i,x in enumerate(d['cables'],5):
    CB.cell(i,1,i-4)
    for j,k in enumerate(['dau','cuoi','cs_dau','cs_cuoi','ma_cap','so_FO','dai_m','ghi_chu'],2):
        v=x[k];CB.cell(i,j,model._f(v) if k in('cs_dau','cs_cuoi','so_FO','dai_m') and model._f(v) is not None else v)
cl=4+len(d['cables']);CB.cell(cl+1,7,'Tổng (m)').font=Font(bold=True);CB.cell(cl+1,8,f'=SUM(H5:H{cl})').font=Font(bold=True)
CB.cell(cl+2,7,'Tổng (km)');CB.cell(cl+2,8,f'=H{cl+1}/1000').number_format='0.0'
box(CB,5,cl,1,9);CB.freeze_panes='A5';CB.auto_filter.ref=f'A4:I{cl}'

# ---------------- Măng xông ----------------
M=sheet('Măng xông')
title(M,'Măng xông (MX) — danh bạ','Sơ đồ hàn chi tiết từng khay chưa nhập: xem sheet "Hàn nối" để điền. MX_NVT_03A/03B là hai măng xông hai bên lề (file ghi nhầm trùng tên).')
header(M,4,['Tên MX','Loại','Vĩ độ','Kinh độ','Mô tả (ghi chú gốc)','Bản đồ'],[26,12,11,11,90,14])
for i,x in enumerate(d['closures'],5):
    for j,v in enumerate([x['name'],'MX thẳng' if x['straight'] else 'MX phối/rẽ',x['lat'],x['lon'],x['desc']],1):M.cell(i,j,v)
    M.cell(i,6,f'=HYPERLINK("https://maps.google.com/?q="&C{i}&","&D{i},"Mở bản đồ")').font=Font(color='0563C1',underline='single')
ML=4+len(d['closures']);box(M,5,ML,1,6);M.freeze_panes='A5'

# ---------------- Hàn nối ----------------
H=sheet('Hàn nối bổ sung')
title(H,'Nhập thêm hàn nối cho măng xông chưa có nhãn (NVT_01, NB-02/03/04, HTP-03B…)','Quy ước nhãn: cùng nhãn = cùng nhóm hàn; trong một nhãn sợi nhỏ hàn sợi nhỏ (A1–A1, A2–A2…). Điền vào ô vàng; ví dụ dưới là DỮ LIỆU GIẢ ĐỊNH để minh hoạ.')
header(H,4,['Măng xông','Cáp A','Sợi A','Nhãn','Cáp B','Sợi B','Khay','Ghi chú'],[22,26,8,8,26,8,8,50])
dv=DataValidation(type='list',formula1=f"='Măng xông'!$A$5:$A${ML}",allow_blank=True);H.add_data_validation(dv);dv.add('A5:A300')
for r in range(5,301):
    for c in range(1,9):H.cell(r,c).fill=INP if r>20 else PatternFill()
ex_ends={'Cáp 12FO (A)':['A','A','B','B','C','C'],'Cáp 6FO (B)':['A','A','B','B']}
pairs,warn=model.splice_pairs(ex_ends)
for i,(lab,a,x,b,y) in enumerate(pairs,5):
    for j,v in enumerate(['(VÍ DỤ GIẢ ĐỊNH)',a,x,lab,b,y,'','minh hoạ quy ước nhãn'],1):H.cell(i,j,v);H.cell(i,j).font=Font(italic=True,color='7F7F7F')
H.cell(5+len(pairs),1,'(VÍ DỤ) cảnh báo: '+'; '.join(warn)).font=Font(italic=True,color='C00000')
H.freeze_panes='A5'
# đường đi theo lề
H2=sheet('Tuyến lề NVT')
title(H2,'Đường đi hai lề trục Nguyễn Văn Tạo (người dùng cung cấp)','Cáp trục lề trái từ NBE12 qua MX_NVT_01 → 02A → 03A là một cáp dùng chung của R2 và R1 (mỗi link 2 core riêng).')
header(H2,4,['Ring','Thứ tự','Điểm','Ghi chú'],[8,8,26,70])
path=[('R2',['AGG-NBE12 (cáp trục lề trái)','MX_NVT_01','HNB054','MX_NVT_01','MX_NVT_02A','HNB006','MX_NVT_02A','lề trái','MX_NVT_03A','HNB091','MX_NVT_03B','lề phải','MX','AGG-NBE22']),
      ('R1',['AGG-NBE12','lề trái','MX_NVT_03A','lề trái','HNB059','lề trái','HNB060','lề phải','HNB048','lề phải','MX_NVT_03B','AGG-NBE22'])]
r=5
for rid,pts_ in path:
    for k,p in enumerate(pts_,1):
        H2.cell(r,1,rid).fill=PatternFill('solid',fgColor=RING_COL[rid]);H2.cell(r,2,k);H2.cell(r,3,p)
        if p.startswith('lề'):H2.cell(r,3).font=Font(italic=True,color='C00000')
        r+=1
    r+=1

# ---------------- Sự cố ----------------
I=sheet('Sự cố đứt cáp')
title(I,'Giả lập đứt cáp — một đoạn bất kỳ','Điều cần tránh: một CSG không còn link nào về AGG. Suy hao các đoạn còn lại không đổi (tính theo từng đoạn).')
header(I,4,['Đoạn đứt','Ring','Từ → Đến','Lề','CSG mất hết link','CSG còn một hướng','Nhận xét'],[10,7,20,14,24,40,50])
r=5
for s in segs:
    res=model.simulate_cut(d,[s['id']]);lost=[k for k,v in res.items() if v==(None,None)];one=[f"{k}→{[x for x in v if x][0]}" for k,v in res.items() if v!=(None,None) and None in v and d['nodes'][k].get('ring')==s['ring']]
    note='trạm nhánh cụt mất kết nối' if s['kind']!='ring' else ('còn bảo vệ: mọi CSG vẫn về được AGG' if not lost else 'MẤT KẾT NỐI')
    for j,v in enumerate([s['id'],s['ring'],f"{s['a']} → {s['z']}",s['roadside'],', '.join(lost) or '—',', '.join(one) or '—',note],1):I.cell(r,j,v)
    if lost:I.cell(r,5).fill=BAD
    r+=1
box(I,5,r-1,1,7)
r+=1;I.cell(r,1,'Cặp đoạn cùng ring đứt cùng lúc làm CSG mất hết link (nguy cơ nếu hai đoạn đi chung cáp/hành lang)').font=Font(bold=True,size=12);r+=1
header(I,r,['Đoạn 1','Ring','Đoạn 2','Chung hành lang (km)','CSG bị cô lập','','Cảnh báo'],None);r+=1;r0=r
for rid in d['rings']:
    rs=sorted([s for s in segs if s['ring']==rid and s['kind']=='ring'],key=lambda s:s['seq'])
    for a,b in itertools.combinations(rs,2):
        res=model.simulate_cut(d,[a['id'],b['id']]);lost=[k for k,v in res.items() if v==(None,None) and d['nodes'][k].get('ring')==rid]
        if not lost:continue
        ov=round(model.corridor_overlap(a['polyline'],b['polyline']),2)
        for j,v in enumerate([a['id'],rid,b['id'],ov,', '.join(lost),'','CẦN KIỂM TRA: hai đoạn chung hành lang' if ov>0.15 else 'hai đoạn riêng biệt'],1):I.cell(r,j,v)
        if ov>0.15:I.cell(r,7).fill=WARN
        r+=1
box(I,r0,r-1,1,7)

# ---------------- Cảnh báo dữ liệu ----------------
W=sheet('Dữ liệu cần xác nhận')
title(W,'Mâu thuẫn / bất thường trong các file nguồn (tự động phát hiện + ghi chú)')
header(W,3,['Đối tượng','Loại','Nội dung','Người xác nhận','Kết quả'],[26,14,90,16,30])
for i,x in enumerate(d['issues'],4):
    for j,v in enumerate(x,1):W.cell(i,j,v)
    W.cell(i,4).fill=INP;W.cell(i,5).fill=INP
box(W,4,3+len(d['issues']),1,5);W.freeze_panes='A4'

# ---------------- Đối tác ----------------
PT=sheet('Đối tác dùng core trống')
title(PT,'Core trống cho đối tác (ghi chú DWDM / VNN / TTC trong ODF AGG)','Chưa có số core cụ thể: điền từ ODF AGG và bảng hàn.')
header(PT,3,['Đối tác','Vị trí/ODF','Ghi chú','Đoạn','Số core'],[24,48,50,16,10])
for i,x in enumerate(d['partners'],4):
    for j,v in enumerate(x,1):PT.cell(i,j,v)
    PT.cell(i,4).fill=INP;PT.cell(i,5).fill=INP
box(PT,4,3+len(d['partners']),1,5)

# ---------------- Vật tư ----------------
V=sheet('Vật tư')
title(V,'Vật tư măng xông / ODF theo sổ ghi tay','ODF theo dung lượng (6/12/24) và danh sách MX thiết kế / phát sinh.')
header(V,3,['Loại','STT','Tên MX','Đầu cáp','Thiết kế/Phát sinh (mối hàn)','Ghi chú'],[14,6,34,10,18,40])
for i,x in enumerate(d['vat_tu'],4):
    for j,k in enumerate(['loai','stt','ten_mx','dau_cap','thiet_ke_hoac_phat_sinh','ghi_chu'],1):
        v=x[k];V.cell(i,j,model._f(v) if k in('dau_cap','thiet_ke_hoac_phat_sinh') and model._f(v) is not None else v)
vl=3+len(d['vat_tu']);V.cell(vl+2,3,'Tổng mối hàn thiết kế').font=Font(bold=True);V.cell(vl+2,5,f'=SUMIF(A4:A{vl},"mx_thiet_ke",E4:E{vl})')
V.cell(vl+3,3,'Tổng mối hàn phát sinh').font=Font(bold=True);V.cell(vl+3,5,f'=SUMIF(A4:A{vl},"mx_phat_sinh",E4:E{vl})')
V.cell(vl+5,3,'ODF theo dung lượng: 6FO × 2 (HNB018, HNB113), 12FO × 27, 24FO × 4 (AGG12, AGG22) — theo sổ ghi tay')
box(V,4,vl,1,6)

# ---------------- Tra cứu ----------------
Q=sheet('Tra cứu')
title(Q,'Tra cứu theo trạm','Chọn trạm ở ô vàng.')
Q['A4']='Trạm';Q['B4']='HNB066';Q['B4'].fill=INP;Q['A4'].font=Font(bold=True)
dv2=DataValidation(type='list',formula1=f"='Nút'!$A$5:$A${NL}");Q.add_data_validation(dv2);dv2.add('B4')
NR=f"'Nút'!"
items=[('Ring',f'=INDEX({NR}C5:C{NL},MATCH(B4,{NR}A5:A{NL},0))'),('Vai trò',f'=INDEX({NR}D5:D{NL},MATCH(B4,{NR}A5:A{NL},0))'),
       ('Vĩ độ',f'=INDEX({NR}E5:E{NL},MATCH(B4,{NR}A5:A{NL},0))'),('Kinh độ',f'=INDEX({NR}F5:F{NL},MATCH(B4,{NR}A5:A{NL},0))'),
       ('Đoạn đi tới trạm (từ phía NBE12)',f'=IFERROR(INDEX({rng("A")},MATCH(B4,{rng("D")},0)),"—")'),
       ('  ← từ',f'=IFERROR(INDEX({rng("C")},MATCH(B4,{rng("D")},0)),"—")'),('  km tính',f'=IFERROR(INDEX({rng("K")},MATCH(B4,{rng("D")},0)),"—")'),
       ('  dB lớn nhất',f'=IFERROR(INDEX({rng("N")},MATCH(B4,{rng("D")},0)),"—")'),('  trạng thái',f'=IFERROR(INDEX({rng("R")},MATCH(B4,{rng("D")},0)),"—")'),
       ('Đoạn đi ra (phía NBE22)',f'=IFERROR(INDEX({rng("A")},MATCH(B4,{rng("C")},0)),"—")'),
       ('  → đến',f'=IFERROR(INDEX({rng("D")},MATCH(B4,{rng("C")},0)),"—")'),('  km tính',f'=IFERROR(INDEX({rng("K")},MATCH(B4,{rng("C")},0)),"—")'),
       ('  dB lớn nhất',f'=IFERROR(INDEX({rng("N")},MATCH(B4,{rng("C")},0)),"—")'),('  trạng thái',f'=IFERROR(INDEX({rng("R")},MATCH(B4,{rng("C")},0)),"—")'),
       ('Số cổng ODF có ghi hướng',f"=COUNTIFS('ODF chi tiết'!A:A,B4,'ODF chi tiết'!D:D,\"<>\")")]
for i,(k,f) in enumerate(items,6):
    Q.cell(i,1,k);Q.cell(i,2,f)
Q.column_dimensions['A'].width=34;Q.column_dimensions['B'].width=18;box(Q,6,5+len(items),1,2)

# ---------------- Sự cố & khắc phục ----------------
X=sheet('Khắc phục sự cố')
title(X,'Triệu chứng → nguyên nhân thường gặp → xử lý (tham khảo, bổ sung theo thực tế)')
header(X,3,['Triệu chứng','Nguyên nhân thường gặp','Cách xử lý / kiểm tra','Dùng sheet'],[38,50,70,22])
tips=[('CSG mất link một hướng, ring vẫn chạy','Đứt/hỏng một đoạn kề','OTDR từ đầu trạm, đổi khoảng cách ra vị trí (MX gần nhất); báo thi công; theo dõi hướng còn lại','Sự cố đứt cáp'),
('CSG mất cả hai hướng','Hai đoạn cùng ring đứt cùng lúc (chung hành lang) hoặc trạm nhánh cụt','Kiểm tra cặp đoạn ở sheet Sự cố; ưu tiên khôi phục đoạn có CSG cô lập','Sự cố đứt cáp'),
('Suy hao một đoạn cao bất thường','Uốn cong, mối hàn xấu, đầu nối bẩn','OTDR hai chiều; vệ sinh đầu nối ODF; kiểm tra khay hàn MX gần nhất','Đoạn & quỹ suy hao'),
('Hai core lệch suy hao > 2 dB','Một core có mối hàn/đầu nối xấu','So sánh OTDR từng core; ví dụ đoạn R2-02 (HNB054→HNB006) chênh 2,02 dB','Đoạn & quỹ suy hao'),
('Chiều dài OTDR khác hẳn hợp đồng/KML','Đổi tuyến, dự trữ lớn, hoặc số liệu ghi sai','Đối chiếu KML & sổ cáp; cập nhật tuyến thực tế (HNB066, HNB005 đã đổi tuyến)','Dữ liệu cần xác nhận'),
('Không tìm thấy điểm đứt theo OTDR','Chỉ số mét (cs) hai đầu cáp thiếu','Bổ sung chỉ số mét hai đầu từng cáp ở sheet Cáp','Cáp'),
('Cổng ODF ghi "hư"','Đầu nối/pigtail hỏng','Chuyển sang cổng dự phòng cùng nhóm, cập nhật trạng thái','ODF AGG')]
for i,x in enumerate(tips,4):
    for j,v in enumerate(x,1):X.cell(i,j,v).alignment=Alignment(wrap_text=True,vertical='top')
box(X,4,3+len(tips),1,4)


# ---------------- Phối core (trọng tâm) ----------------
cabinfo={c['cap_id']:c for c in d['mx_cables']}
def cname(i):
    c=cabinfo[i];h=c['huong'] or c['tieu_de']
    return f"{c['o_goc']} · {(c['ma_cap']+' ') if c['ma_cap'] else ''}{c['so_FO']}FO → {h}"
PC=sheet('Phối core MX')
title(PC,'Bảng phối core tại măng xông: core nào hàn với core nào','Suy từ nhãn trong sơ đồ MX (cùng nhãn = cùng nhóm hàn; core nhỏ hàn core nhỏ). Màu sợi/ống theo thứ tự 6 sợi/ống (giả định). Cột "Trạm" lấy từ tên trạm ghi sát nhãn.')
header(PC,4,['Măng xông','Nhãn','Cáp A','Core A','Ống A','Màu A','Cáp B','Core B','Ống B','Màu B','Trạm (ghi trong sơ đồ)','Ghi chú'],[16,7,46,7,6,11,46,7,6,11,16,30])
for i,p in enumerate(d['mx_pairs'],5):
    for j,v in enumerate([p['mx'],p['label'],cname(p['a']),p['core_a'],int(p['tube_a']),p['color_a'],cname(p['b']),p['core_b'],int(p['tube_b']),p['color_b'],p['station'],''],1):PC.cell(i,j,v)
    PC.cell(i,12).fill=INP
pl=4+len(d['mx_pairs']);box(PC,5,pl,1,12);PC.freeze_panes='C5';PC.auto_filter.ref=f'A4:L{pl}'
PCA="'Phối core MX'!"
CM=sheet('Cáp tại MX')
title(CM,'Cáp (đầu cáp) tại từng măng xông','Ô gốc = vị trí ô trong sơ đồ MX. Số core vẽ / có nhãn tính từ sheet "Core MX".')
header(CM,4,['Măng xông','Ô gốc','Mã cáp','Loại','Số FO','Chỉ số mét','Hướng','Số core vẽ','Số core có nhãn','Tiêu đề gốc'],[16,8,18,8,7,10,34,9,10,60])
for i,c in enumerate(d['mx_cables'],5):
    for j,v in enumerate([c['mx'],c['o_goc'],c['ma_cap'],c['loai'],model._f(c['so_FO']),model._f(c['chi_so_m']),c['huong']],1):CM.cell(i,j,v)
    CM.cell(i,8,f"=COUNTIFS('Core MX'!B:B,\"{c['cap_id']}\")");CM.cell(i,9,f"=COUNTIFS('Core MX'!B:B,\"{c['cap_id']}\",'Core MX'!F:F,\"?*\")");CM.cell(i,10,c['tieu_de'])
cl2=4+len(d['mx_cables']);box(CM,5,cl2,1,10);CM.freeze_panes='C5';CM.auto_filter.ref=f'A4:J{cl2}'
CO=sheet('Core MX')
title(CO,'Từng core của từng đầu cáp tại măng xông','Nhãn trống = chưa ghi (dự phòng/chưa phối).')
header(CO,4,['Măng xông','Mã đầu cáp','Core','Ống','Màu sợi','Nhãn','Trạm'],[16,26,6,6,11,7,12])
for i,r in enumerate(d['mx_cores'],5):
    for j,v in enumerate([r['mx'],r['cap_id'],int(r['core']),int(r['ong']),r['mau_soi'],r['nhan'],r['tram']],1):CO.cell(i,j,v)
box(CO,5,4+len(d['mx_cores']),1,7);CO.freeze_panes='A5';CO.auto_filter.ref=f'A4:G{4+len(d["mx_cores"])}'
CT=sheet('Core theo trạm')
title(CT,'Core đi vào từng trạm theo sơ đồ MX','Gom theo tên trạm ghi trong sơ đồ: măng xông, cáp, nhóm core và nhãn.')
header(CT,4,['Trạm','Măng xông','Nhãn','Cáp','Các core','Số core'],[14,16,7,56,24,8])
grp={}
for r in d['mx_cores']:
    if r['tram']:grp.setdefault((r['tram'],r['mx'],r['nhan'],r['cap_id']),[]).append(int(r['core']))
for i,((tr,mx,lab,cid),cs_) in enumerate(sorted(grp.items()),5):
    for j,v in enumerate([tr,mx,lab,cname(cid),', '.join(map(str,cs_)),len(cs_)],1):CT.cell(i,j,v)
box(CT,5,4+len(grp),1,6);CT.freeze_panes='A5';CT.auto_filter.ref=f'A4:F{4+len(grp)}'
CW_=sheet('Cảnh báo phối core')
title(CW_,'Cảnh báo khi suy bảng phối core từ sơ đồ MX')
header(CW_,3,['Măng xông','Nội dung','Xác nhận'],[18,100,30])
rr=4
for mx,wmsg in d['mx_warn']:
    CW_.cell(rr,1,mx);CW_.cell(rr,2,wmsg);CW_.cell(rr,3).fill=INP;rr+=1
for o,t in d['mx_notes']:
    CW_.cell(rr,1,'ghi chú tay '+o);CW_.cell(rr,2,t);CW_.cell(rr,3).fill=INP;rr+=1
box(CW_,4,rr-1,1,3)

# ---------------- Đường đi link ----------------
PA=sheet('Đường đi link')
title(PA,'Đường đi từng link ring qua ODF – cáp – măng xông (SFP 1310 nm, 2 core/link)','Mô tả do người dùng cung cấp; độ dài cáp từ sổ độ dài; đối chiếu nhãn với sơ đồ MX. Link chưa hoàn chỉnh ghi rõ.')
header(PA,4,['Link','Từ','Đến','Tổng cáp (m)','OTDR tham chiếu (m)','Lệch (m)','Lệch %','Kết luận','Ghi chú'],[9,9,9,12,14,10,9,26,60])
nl_=len(d['paths']);r0=5
for i,lk in enumerate(d['paths']):
    r=r0+i;ref=(lk['odf_otdr_m'] or [None])[0] or lk['bbnt_otdr_m']
    PA.cell(r,1,lk['id']);PA.cell(r,2,lk['frm']);PA.cell(r,3,lk['to'])
    PA.cell(r,4,f'=SUMIFS($G${r0+nl_+4}:$G$2000,$A${r0+nl_+4}:$A$2000,A{r},$C${r0+nl_+4}:$C$2000,"cáp")');PA.cell(r,5,ref)
    PA.cell(r,6,f'=D{r}-E{r}');PA.cell(r,7,f'=F{r}/E{r}');PA.cell(r,7).number_format='0.0%'
    PA.cell(r,8,('Chưa hoàn chỉnh (còn tiếp)' if not lk['complete'] else f'=IF(ABS(G{r})<=0.02,"Khớp (≤2%)","Lệch >2%: kiểm tra")'))
    PA.cell(r,9,'OTDR tham chiếu: '+('ODF AGG ghi 982/988 m' if lk['odf_otdr_m'] else 'BBNT (số thống nhất)'))
box(PA,r0,r0+nl_-1,1,9)
hr=r0+nl_+2;PA.cell(hr,1,'Chi tiết từng bước').font=Font(bold=True,size=12)
header(PA,hr+1,['Link','Bước','Loại','Vị trí / cáp','Từ → Đến','Core','Dài (m)','Cộng dồn (m)','Đối chiếu sơ đồ MX','Ghi chú'],None)
PA.column_dimensions['J'].width=50
rr=hr+2
for lk in d['paths']:
    first=rr
    for k,s in enumerate(lk['steps'],1):
        PA.cell(rr,1,lk['id']);PA.cell(rr,2,k)
        if s['t']=='odf':
            vals=['ODF',f"{s['at']}: {s['name']}",'', '+'.join(map(str,s['cores'])),None,None,'',s.get('note','')]
        elif s['t']=='cable':
            vals=['cáp',s['name'],f"{s['frm']} → {s['to']}",'+'.join(map(str,s['cores'])),s['len_m'],f'=SUM(G{first}:G{rr})' ,'',s['src']]
        elif s['t']=='mx':
            vals=['măng xông',s['at'],'',f"{'+'.join(map(str,s['cores_a']))} ↔ {'+'.join(map(str,s['cores_b']))}",None,None,f"{s['check']} — {s['check_detail']}",s.get('note','')]
        else:
            vals=['…','','','',None,None,'',s['note']]
        for j,v in enumerate(vals,3):PA.cell(rr,j,v)
        PA.cell(rr,3).font=Font(bold=True)
        c=PA.cell(rr,9);t=str(vals[6])
        if t.startswith('khớp sơ đồ'):c.fill=GOOD
        elif t.startswith('KHÁC') or 'THỨ TỰ' in t:c.fill=BAD
        elif t.startswith('chưa'):c.fill=WARN
        if s['t']=='mx':PA.cell(rr,3).fill=GREY
        if s['t']=='todo':PA.cell(rr,4).fill=INP
        rr+=1
box(PA,hr+2,rr-1,1,10);PA.freeze_panes='A5'

# ---------------- Tổng quan ----------------
T=wb.create_sheet('Tổng quan',0);T.sheet_view.showGridLines=False
title(T,'Mạng ring CSG Nhà Bè — tổng quan','Dữ liệu từ file ghi tay (BBNT, ODF, MX, độ dài cáp, KML). Công thức sống; sửa Tham số để tính lại.')
kp=[('Số ring',f'=COUNTA(Ring!A5:A12)'),('Số CSG ring',f'=SUM(Ring!C5:C12)'),('Số trạm nhánh cụt (cấp 3)','=COUNTIF(Nút!D5:D%d,"nhánh cụt*")'%NL),
    ('Số đoạn cáp (gồm nhánh cụt)',f'=COUNTA({rng("A")})'),('Tổng km tính (đoạn ring)',f'=SUMIF({rng("E")},"ring",{rng("K")})'),
    ('Đoạn dài nhất (km)',f'=MAX({rng("K")})'),('Dư suy hao thấp nhất (dB)',f'=MIN({rng("P")})'),
    ('Số đoạn Đạt',f'=COUNTIF({rng("R")},"Đạt")'),('Số đoạn Cảnh báo',f'=COUNTIF({rng("R")},"Cảnh báo")'),('Số đoạn chưa đo',f'=COUNTIF({rng("R")},"Chưa đo")'),
    ('Số măng xông trong danh bạ',f"=COUNTA('Măng xông'!A5:A{ML})"),('Số cổng ODF CSG ghi hướng',"=COUNTIFS('ODF chi tiết'!D:D,\"<>\")-1"),('Cặp core hàn (suy từ nhãn MX)',f"=COUNTA('Phối core MX'!A5:A{pl})"),('Măng xông đã có nhãn core',f"=SUMPRODUCT(1/COUNTIF('Phối core MX'!A5:A{pl},'Phối core MX'!A5:A{pl}))"),('Mục cần xác nhận',f"=COUNTA('Dữ liệu cần xác nhận'!A4:A{3+len(d['issues'])})")]
header(T,4,['Chỉ số','Giá trị'],[44,16])
for i,(k,f) in enumerate(kp,5):
    T.cell(i,1,k);x=T.cell(i,2,f);x.font=Font(bold=True,size=12);x.number_format='0.0' if 'km' in k or 'dB' in k else '0'
box(T,5,4+len(kp),1,2)
T.conditional_formatting.add('B13',CellIsRule(operator='greaterThan',formula=['0'],fill=WARN))
notes=['Cách đọc: sheet "Ring" tóm tắt 8 ring; "Đoạn & quỹ suy hao" là bảng trung tâm; "Sự cố đứt cáp" liệt kê hệ quả đứt từng đoạn.',
 'Quy ước: bảo vệ trong phạm vi từng ring AGG–AGG; các ring được phép chung cáp; nguy cơ cần tránh là một CSG không còn link nào về AGG.',
 'Suy hao tính theo từng đoạn giữa hai trạm liền kề (SFP thu/phát mỗi trạm), ngưỡng 12,1 / 17 / 27 dB theo chiều dài; lệch hai core ≤ 2 dB.',
 'Chưa có: bảng hàn từng khay của măng xông (chỉ có danh bạ), số core cụ thể của đối tác, chỉ số mét hai đầu cáp. Đã đánh dấu để bổ sung.']
for i,t in enumerate(notes,20):T.cell(i,1,t);T.cell(i,1).alignment=Alignment(wrap_text=True)
T.merge_cells('A20:H20');T.merge_cells('A21:H21');T.merge_cells('A22:H22');T.merge_cells('A23:H23')
for r in range(20,24):T.row_dimensions[r].height=32
for c in 'CDEFGH':T.column_dimensions[c].width=14

order=['Tổng quan','Phối core MX','Đường đi link','Core theo trạm','Cáp tại MX','Core MX','Cảnh báo phối core','Ring','Đoạn & quỹ suy hao','Sự cố đứt cáp','Nút','ODF CSG','ODF chi tiết','ODF AGG','Cáp','Măng xông','Hàn nối bổ sung','Tuyến lề NVT','Đối tác dùng core trống','Dữ liệu cần xác nhận','Khắc phục sự cố','Vật tư','Tra cứu','Tham số']
wb._sheets=[wb[n] for n in order]
for ws in wb:
    ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0;ws.sheet_properties.pageSetUpPr.fitToPage=True
out=os.path.join(os.path.dirname(os.path.abspath(__file__)),'ring_network_NhaBe.xlsx');wb.save(out);print('saved',out)
